"""Convert one document to Markdown meant as context for AI models.

Design rules:
- Nothing visible in the original may silently disappear. Every PDF page is
  checked against its raw text layer; image text is recognised by OCR.
- Human and AI must be able to point at the same place: PDF pages and slides
  carry visible markers with the page number shown by PDF viewers (plus the
  printed page label when it differs), spreadsheets keep Excel cell addresses.
- Layout only matters where it carries meaning (headings, lists, tables).
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime, time
from pathlib import Path

import ocr

IMAGES = {".png", ".jpg", ".jpeg", ".gif", ".tif", ".tiff", ".heic", ".bmp"}
SHEETS = {".xlsx", ".xlsm", ".xls", ".csv"}
SUPPORTED = {".pdf", ".docx", ".pptx", ".html", ".htm"} | SHEETS | IMAGES

# Images smaller than this share of the page (logos, icons) are not worth OCR.
MIN_IMAGE_SHARE = 0.03
# Pages with less text than this and a large image are scans and recognised as a whole.
SCAN_TEXT_LIMIT = 50
# Share of the raw text-layer words that must survive layout analysis, otherwise the raw text is used.
MIN_RECALL = 0.9
# Spreadsheets: beyond this, tables stop being useful context and are cut (with a note).
MAX_SHEET_ROWS = 3000
MAX_FORMULAS = 400


class ConversionError(Exception):
    """Expected problem with a source file; the message is shown to the user."""


# --- shared helpers -------------------------------------------------------------

def _words(text: str) -> set[str]:
    return set(re.findall(r"\w{2,}", unicodedata.normalize("NFKC", text).casefold()))


def _garbled(text: str) -> bool:
    """Broken font encodings yield private-use or replacement characters instead of letters."""
    letters = [c for c in text if not c.isspace()]
    if len(letters) < 20:
        return False
    bad = sum(1 for c in letters if c == "\ufffd" or unicodedata.category(c) in ("Co", "Cn", "Cc"))
    return bad / len(letters) > 0.3


def _new_ocr_lines(recognised: str, known: set[str]) -> str:
    """Drop OCR lines that only repeat text already present (e.g. text drawn over a background image)."""
    kept = []
    for line in recognised.splitlines():
        words = _words(line)
        if words and len(words & known) / len(words) >= 0.8:
            continue
        kept.append(line)
    return "\n".join(kept).strip()


def _clean(text: str) -> str:
    """Remove download-portal watermarks that carry no course content."""
    text = re.sub(r"^\s*lOMoARcPSD\|\d+\s*$", "", text, flags=re.M)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _figure(text: str, title: str = "Abbildung") -> str:
    body = "\n".join(f"> {line}" if line else ">" for line in text.splitlines())
    return f"> [{title} – Text per OCR]\n{body}\n"


# --- PDF ------------------------------------------------------------------------

# The layout model needs ~0.1 s for a normal slide but can take 40 s on poster-sized pages
# with thousands of vector paths and large images. Such documents are analysed in a child
# process with a time budget; pages it does not finish fall back to the raw text layer.
LAYOUT_SECONDS_PER_PAGE = 4.0
HEAVY_BYTES_PER_PAGE = 400_000
A4_AREA = 595 * 842


def _layout_batches(path: str, page_count: int, queue) -> None:
    """Child process: send {page: markdown} per batch; small batches keep losses small on timeout."""
    import pymupdf
    import pymupdf4llm

    doc = pymupdf.open(path)
    for start in range(0, page_count, 3):
        pages = list(range(start, min(start + 3, page_count)))
        chunks = pymupdf4llm.to_markdown(doc, pages=pages, page_chunks=True, show_progress=False, use_ocr=False)
        queue.put({c["metadata"]["page_number"]: c["text"].strip() for c in chunks})
    queue.put(None)


def _layout_pages(doc, source: Path) -> tuple[dict[int, str], int]:
    """Markdown per page from pymupdf4llm's layout analysis.

    Returns {page number: markdown} and the number of leading pages analysed.
    """
    try:
        import pymupdf4llm
    except ImportError:
        return {}, 0
    heavy = (source.stat().st_size / doc.page_count > HEAVY_BYTES_PER_PAGE
             or max(p.rect.width * p.rect.height for p in doc) > 2.1 * A4_AREA)
    if not heavy:
        try:
            # use_ocr=False: our own OCR (macOS Vision) handles images and scans.
            chunks = pymupdf4llm.to_markdown(doc, page_chunks=True, show_progress=False, use_ocr=False)
        except Exception:  # best effort; the raw text layer is the safety net
            return {}, 0
        return {c["metadata"]["page_number"]: c["text"].strip() for c in chunks}, doc.page_count

    import multiprocessing
    import queue as queues

    context = multiprocessing.get_context("spawn")
    channel = context.Queue()
    child = context.Process(target=_layout_batches, args=(str(source), doc.page_count, channel), daemon=True)
    child.start()
    layout: dict[int, str] = {}
    try:
        # The first batch also pays for starting Python and loading the layout model.
        timeout = 20 + 3 * LAYOUT_SECONDS_PER_PAGE
        while True:
            batch = channel.get(timeout=timeout)
            if batch is None:
                break
            layout.update(batch)
            timeout = 3 * LAYOUT_SECONDS_PER_PAGE
    except queues.Empty:
        pass
    finally:
        if child.is_alive():
            child.kill()
        child.join(1)
    done = 0
    while done + 1 in layout:
        done += 1
    return layout, done


def pdf(source: Path) -> tuple[str, dict]:
    import pymupdf

    try:
        doc = pymupdf.open(source)
    except pymupdf.EmptyFileError as exc:
        raise ConversionError("Datei ist leer (0 Byte)") from exc
    except (pymupdf.FileDataError, RuntimeError) as exc:
        raise ConversionError("PDF ist beschädigt oder keine PDF-Datei") from exc
    if doc.needs_pass:
        raise ConversionError("PDF ist passwortgeschützt und ohne Passwort nicht lesbar")
    if doc.page_count == 0:
        raise ConversionError("PDF enthält keine Seiten")

    layout, layout_until = _layout_pages(doc, source)

    usage: dict[int, int] = {}
    for page in doc:
        for xref in {info[0] for info in page.get_images()}:
            usage[xref] = usage.get(xref, 0) + 1
    # Images on most pages are logos or slide templates.
    decoration = {x for x, n in usage.items() if doc.page_count >= 4 and n > doc.page_count / 2}
    cache: dict[int, str] = {}
    stats = {"seiten": doc.page_count, "ocr_seiten": [], "rohtext_seiten": [], "ohne_text": []}
    if layout_until < doc.page_count:
        stats["layout_bis"] = layout_until
    parts = []
    for index, page in enumerate(doc):
        number = index + 1
        label = page.get_label()
        marker = f"[Seite {number} von {doc.page_count}"
        if label and label != str(number):
            marker += f" · im Dokument: {label}"
        marker += "]"

        # Content-stream order: unlike sort=True it keeps overprinted "fake bold" text intact.
        raw = _clean(page.get_text("text"))
        text = _clean(layout.get(number, ""))
        raw_words = _words(raw)
        lost = raw_words and len(raw_words & _words(text)) / len(raw_words) < MIN_RECALL
        # Rotated text sometimes comes back as one word per line.
        shredded = text.count("\n") > 3 * (raw.count("\n") + 1) + 5
        if lost or shredded:
            text = raw
            stats["rohtext_seiten"].append(number)
        area = page.rect.width * page.rect.height
        images, places = [], set()
        for info in page.get_image_info(xrefs=True):
            rect = pymupdf.Rect(info["bbox"])
            place = tuple(round(v) for v in rect)
            # The same picture is often listed twice (with and without xref).
            if rect.get_area() >= MIN_IMAGE_SHARE * area and place not in places:
                places.add(place)
                images.append(info)

        sparse = len(raw) < SCAN_TEXT_LIMIT
        # No usable text layer but something is drawn: a scan, or text converted to outlines.
        if (sparse and (images or len(page.get_drawings()) > 20)) or _garbled(raw):
            recognised = ocr.recognize(page.get_pixmap(dpi=200).tobytes("png"))
            if recognised:
                stats["ocr_seiten"].append(number)
                note = "Textebene unlesbar" if not sparse else "Seite ohne Textebene"
                extra = _new_ocr_lines(raw, _words(recognised)) if not _garbled(raw) else ""
                parts.append(f"{marker}\n\n[{note} – gesamter Text per OCR]\n\n{recognised}\n"
                             + (f"\n{extra}\n" if extra else ""))
                continue

        known = _words(text)
        figures = []
        for info in images:
            xref = info.get("xref", 0)
            if xref in decoration:
                continue
            if xref not in cache or not xref:
                rect = pymupdf.Rect(info["bbox"]) & page.rect
                if rect.is_empty:
                    continue
                # Render the area (applies masks/rotation) at a resolution Vision reads well.
                scale = min(4.0, max(2.0, 1200 / max(rect.width, rect.height)))
                cache[xref] = ocr.recognize(page.get_pixmap(clip=rect, matrix=pymupdf.Matrix(scale, scale)).tobytes("png"))
            fresh = _new_ocr_lines(cache[xref], known)
            if len(fresh) >= 4:
                figures.append(_figure(fresh))
                known |= _words(fresh)
        if not text and not figures:
            stats["ohne_text"].append(number)
        parts.append(f"{marker}\n\n{text}\n" + "".join(f"\n{f}" for f in figures))
    return "\n".join(parts), stats


# --- Word / PowerPoint / HTML (markitdown) + OCR for embedded pictures -----------

def _markitdown(source: Path) -> str:
    from markitdown import MarkItDown

    try:
        return MarkItDown(enable_plugins=False).convert(str(source)).text_content.strip()
    except Exception as exc:
        if source.stat().st_size == 0:
            raise ConversionError("Datei ist leer (0 Byte)") from exc
        if "zip" in str(exc).lower():
            raise ConversionError("Datei ist beschädigt oder kein gültiges Office-Dokument") from exc
        raise


def _office_images(source: Path) -> dict[int, list[bytes]]:
    """Embedded pictures per slide (pptx) or in document order under key 0 (docx)."""
    import posixpath
    import zipfile
    from xml.etree import ElementTree

    found: dict[int, list[bytes]] = {}
    if source.suffix.lower() == ".pptx":
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        def pictures(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from pictures(shape.shapes)
                elif hasattr(shape, "image"):
                    try:
                        yield shape.image.blob, shape.image.ext
                    except (AttributeError, ValueError, KeyError):
                        pass

        for number, slide in enumerate(Presentation(str(source)).slides, 1):
            found[number] = [blob for blob, ext in pictures(slide.shapes) if ext.lower() not in ("wmf", "emf", "svg")]
        return found

    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    with zipfile.ZipFile(source) as archive:
        targets = {rel.get("Id"): rel.get("Target")
                   for rel in ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))}
        document = ElementTree.fromstring(archive.read("word/document.xml"))
        seen = set()
        for blip in document.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
            target = targets.get(blip.get(f"{rel_ns}embed") or "")
            if not target or target in seen or target.lower().endswith((".wmf", ".emf", ".svg")):
                continue
            seen.add(target)
            name = target[1:] if target.startswith("/") else posixpath.normpath(posixpath.join("word", target))
            try:
                found.setdefault(0, []).append(archive.read(name))
            except KeyError:
                pass
    return found


def office(source: Path) -> tuple[str, dict]:
    body = _markitdown(source)
    # Picture placeholders ("![Grafik 4](Grafik4.jpg)", truncated data URIs) carry no content.
    body = re.sub(r"!\[[^\]]*\]\([^)]*\)", "[Abbildung]", body)
    known = _words(body)
    texts = {}
    for key, blobs in _office_images(source).items():
        recognised = [_new_ocr_lines(ocr.recognize(blob), known) for blob in blobs]
        texts[key] = [t for t in recognised if len(t) >= 4]
    stats: dict = {}
    if source.suffix.lower() == ".pptx":
        pieces = re.split(r"<!-- Slide number: (\d+) -->", body)
        total = (len(pieces) - 1) // 2
        stats["folien"] = total
        out = [pieces[0].strip()]
        for i in range(1, len(pieces), 2):
            number = int(pieces[i])
            out.append(f"[Folie {number} von {total}]\n\n{pieces[i + 1].strip()}\n")
            out += [_figure(text) for text in texts.get(number, [])]
        body = "\n".join(p for p in out if p)
    elif texts.get(0):
        body += "\n\n## Text aus Abbildungen (OCR, in Dokumentreihenfolge)\n\n" + "\n".join(
            _figure(text, f"Abbildung {n}") for n, text in enumerate(texts[0], 1))
    return body + "\n", stats


def html(source: Path) -> tuple[str, dict]:
    return _markitdown(source) + "\n", {}


# --- Spreadsheets -------------------------------------------------------------------

def _column(index: int) -> str:
    name = ""
    while index:
        index, rest = divmod(index - 1, 26)
        name = chr(65 + rest) + name
    return name


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "WAHR" if value else "FALSCH"
    if isinstance(value, float):
        value = f"{value:.10g}"
    elif isinstance(value, (datetime, date, time)):
        value = value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).replace("|", "\\|").replace("\r", "").replace("\n", " ⏎ ").strip()


def _sheet_table(rows: list[tuple[int, list]], first_col: int) -> str:
    """Markdown table whose header holds Excel column letters and whose first column holds row numbers."""
    width = max(len(values) for _, values in rows)
    header = "| Zeile | " + " | ".join(_column(first_col + i) for i in range(width)) + " |"
    lines = [header, "|" + " --- |" * (width + 1)]
    for number, values in rows:
        cells = [_cell(v) for v in values] + [""] * (width - len(values))
        lines.append(f"| {number} | " + " | ".join(cells) + " |")
    return "\n".join(lines)


def _trim(grid: dict[tuple[int, int], object]) -> tuple[list[tuple[int, list]], int, int]:
    """Rows of the used range (non-empty rows only); returns rows, first column, dropped rows."""
    if not grid:
        return [], 1, 0
    cols = [c for _, c in grid]
    first, last = min(cols), max(cols)
    rows: dict[int, list] = {}
    for (r, c), value in grid.items():
        rows.setdefault(r, [None] * (last - first + 1))[c - first] = value
    ordered = sorted(rows.items())
    dropped = max(0, len(ordered) - MAX_SHEET_ROWS)
    return ordered[:MAX_SHEET_ROWS], first, dropped


def _sheet_markdown(name: str, hidden: bool, grid: dict, formulas: list[tuple[str, str]]) -> str:
    title = f"## Tabellenblatt „{name}“" + (" (ausgeblendet)" if hidden else "")
    rows, first, dropped = _trim(grid)
    if not rows:
        return f"{title}\n\n(leeres Blatt)\n"
    parts = [title, "", _sheet_table(rows, first)]
    if dropped:
        parts.append(f"\n[… {dropped} weitere Zeilen gekürzt – vollständige Daten im Original]")
    if formulas:
        parts += ["", "Formeln (Zelle: Formel):", ""]
        parts += [f"- {ref}: `{formula}`" for ref, formula in formulas[:MAX_FORMULAS]]
        if len(formulas) > MAX_FORMULAS:
            parts.append(f"- … {len(formulas) - MAX_FORMULAS} weitere Formeln")
    return "\n".join(parts) + "\n"


def spreadsheet(source: Path) -> tuple[str, dict]:
    suffix = source.suffix.lower()
    if source.stat().st_size == 0:
        raise ConversionError("Datei ist leer (0 Byte)")
    sheets = []
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl

        try:
            values = openpyxl.load_workbook(source, data_only=True, read_only=False)
            formulas = openpyxl.load_workbook(source, data_only=False, read_only=False)
        except Exception as exc:
            raise ConversionError("Excel-Datei ist beschädigt oder verschlüsselt") from exc
        for ws in values.worksheets:
            fs = formulas[ws.title]
            grid, found = {}, []
            for row in fs.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    shown = ws.cell(row=cell.row, column=cell.column).value
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        found.append((cell.coordinate, cell.value))
                        # Files saved by tools other than Excel may lack cached results.
                        grid[(cell.row, cell.column)] = shown if shown is not None else cell.value
                    else:
                        grid[(cell.row, cell.column)] = shown
            sheets.append(_sheet_markdown(ws.title, ws.sheet_state != "visible", grid, found))
    elif suffix == ".xls":
        import xlrd

        try:
            # Read normally first: xlrd's mmap fails on iCloud placeholders instead of downloading them.
            book = xlrd.open_workbook(file_contents=source.read_bytes())
        except Exception as exc:
            raise ConversionError("Excel-Datei (.xls) ist beschädigt oder verschlüsselt") from exc
        for ws in book.sheets():
            grid = {}
            for r in range(ws.nrows):
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                        except Exception:
                            pass
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "#FEHLER")
                    grid[(r + 1, c + 1)] = value
            sheets.append(_sheet_markdown(ws.name, ws.visibility != 0, grid, []))
        sheets.append("[Hinweis: Altes .xls-Format – nur Werte, Formeln sind nicht auslesbar.]\n")
    else:  # CSV: guess encoding and delimiter instead of assuming UTF-8 and commas
        raw = source.read_bytes()
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        grid = {(r, c): value for r, row in enumerate(csv.reader(io.StringIO(text), dialect), 1)
                for c, value in enumerate(row, 1) if value != ""}
        sheets.append(_sheet_markdown(source.stem, False, grid, []).replace("## Tabellenblatt", "## Tabelle", 1))
    return "\n".join(sheets), {"blaetter": len(sheets)}


# --- Images -----------------------------------------------------------------------

def image(source: Path) -> tuple[str, dict]:
    if source.stat().st_size == 0:
        raise ConversionError("Datei ist leer (0 Byte)")
    text = ocr.recognize(source.read_bytes())
    return ("[Bilddatei – Text per OCR]\n\n" + text + "\n") if text else "[Bilddatei – kein Text erkannt]\n", {}


def convert(source: Path) -> tuple[str, dict]:
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        return pdf(source)
    if suffix in (".docx", ".pptx"):
        return office(source)
    if suffix in SHEETS:
        return spreadsheet(source)
    if suffix in IMAGES:
        return image(source)
    if suffix in (".html", ".htm"):
        return html(source)
    raise ConversionError(f"Dateityp {suffix} wird nicht unterstützt")
